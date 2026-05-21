from __future__ import annotations

import argparse
import csv
import json
import math
import shutil
import textwrap
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "demo-run" / "deliverables"
ASSET_DIR = OUT_DIR / "assets"
TEMPLATE_BG_GLOB = ROOT / "outputs"

UQ_PURPLE = "#51247A"
UQ_DARK = "#2A1745"
INK = "#14213D"
MUTED = "#4B5563"
LIGHT = "#F5F7FA"
LINE = "#D8DEE9"
TEAL = "#007C89"
CORAL = "#D1442E"
GOLD = "#A56B00"
GREEN = "#2E7D32"
BLUE = "#2F6FDB"

FONT_REGULAR = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
FONT_BOLD = Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")
FONT_BLACK = Path("/System/Library/Fonts/Supplemental/Arial Black.ttf")


def font(size: int, bold: bool = False, black: bool = False, scale: int = 1) -> ImageFont.FreeTypeFont:
    path = FONT_BLACK if black else FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(str(path), size * scale)


def latest_template_background() -> Path | None:
    candidates = sorted(
        TEMPLATE_BG_GLOB.glob("*/presentations/deepfuzz-demo-day/template-inspect/source-slides/source-slide-02.png"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return candidates[0] if candidates else None


def load_json(path: Path) -> dict:
    with path.open() as handle:
        return json.load(handle)


def count_tests(lib: str) -> int:
    return len(list((ROOT / "pipeline_runs" / lib / f"{lib}-thesis-all" / "generated_tests").glob("test_*.py")))


def collect_metrics() -> dict:
    libs = ["jax", "tensorflow", "torch"]
    rows = []
    for lib in libs:
        report = load_json(ROOT / "pipeline_runs" / lib / f"{lib}-thesis-all" / "final_report.json")
        coverage = load_json(ROOT / "stage4" / "results" / f"{lib}-thesis-all-coverage" / "coverage_report.json")
        health = report["pipeline_health"]
        stage4 = health["stage4"]
        rows.append(
            {
                "lib": lib,
                "label": {"jax": "JAX", "tensorflow": "TensorFlow", "torch": "PyTorch"}[lib],
                "version": coverage["target_version"],
                "accepted": health["accepted"]["accepted_apis"],
                "stage2_valid": health["stage2"]["total_valid"],
                "stage3_ready": health["stage3"]["valid"],
                "stage4_eval": stage4["executed_apis"],
                "base_valid": stage4["base_success"],
                "generated_tests": count_tests(lib),
                "api_exec": coverage["api_execution_coverage_percent"],
                "function_cov": coverage["function_line_coverage_percent"],
                "package_cov": coverage["python_line_coverage_percent"],
                "valid_programs": stage4["total_valid_programs"],
                "unique_valid_programs": stage4["unique_valid_programs"],
                "candidate_bugs": len(report["candidate_bugs"]),
                "device_mismatches": stage4["device_oracle_mismatches"],
                "expected_negative": stage4["expected_negative_rejections"],
                "negative_accepted_not_bug": stage4["negative_accepted_not_bug"],
                "pipeline_errors": len(report["pipeline_errors"]),
                "platform_exclusions": len(report["platform_exclusions"]),
                "stage3_failed": health["stage3"]["failed"],
            }
        )

    totals = {}
    for key in [
        "accepted",
        "stage2_valid",
        "stage3_ready",
        "stage4_eval",
        "base_valid",
        "generated_tests",
        "valid_programs",
        "unique_valid_programs",
        "candidate_bugs",
        "device_mismatches",
        "expected_negative",
        "negative_accepted_not_bug",
        "pipeline_errors",
        "platform_exclusions",
        "stage3_failed",
    ]:
        totals[key] = sum(row[key] for row in rows)
    totals["stage2_rate"] = totals["stage2_valid"] / totals["accepted"] * 100
    totals["stage3_rate"] = totals["stage3_ready"] / totals["accepted"] * 100
    totals["base_valid_rate"] = totals["base_valid"] / totals["stage4_eval"] * 100

    return {"rows": rows, "totals": totals}


def fmt_int(value: int) -> str:
    return f"{value:,}"


def fmt_pct(value: float, digits: int = 1) -> str:
    return f"{value:.{digits}f}%"


class Poster:
    def __init__(self, background: Image.Image, scale: int = 3):
        self.scale = scale
        self.logical_w, self.logical_h = background.size
        self.image = background.resize(
            (self.logical_w * scale, self.logical_h * scale),
            Image.Resampling.LANCZOS,
        ).convert("RGB")
        self.draw = ImageDraw.Draw(self.image)

    def xy(self, x: float, y: float) -> tuple[int, int]:
        return int(round(x * self.scale)), int(round(y * self.scale))

    def box(self, x: float, y: float, w: float, h: float) -> tuple[int, int, int, int]:
        return (
            int(round(x * self.scale)),
            int(round(y * self.scale)),
            int(round((x + w) * self.scale)),
            int(round((y + h) * self.scale)),
        )

    def rect(self, x: float, y: float, w: float, h: float, fill: str, outline: str | None = None, width: int = 1, radius: int = 6):
        self.draw.rounded_rectangle(
            self.box(x, y, w, h),
            radius=radius * self.scale,
            fill=fill,
            outline=outline,
            width=max(1, width * self.scale) if outline else 1,
        )

    def line(self, x1: float, y1: float, x2: float, y2: float, fill: str, width: int = 1):
        self.draw.line((*self.xy(x1, y1), *self.xy(x2, y2)), fill=fill, width=max(1, width * self.scale))

    def text(
        self,
        value: str,
        x: float,
        y: float,
        size: int,
        fill: str = INK,
        bold: bool = False,
        black: bool = False,
        anchor: str | None = None,
        align: str = "left",
        max_width: int | None = None,
        line_gap: int = 3,
    ):
        fnt = font(size, bold=bold, black=black, scale=self.scale)
        if max_width:
            value = wrap_text(value, fnt, max_width * self.scale)
        self.draw.multiline_text(
            self.xy(x, y),
            value,
            font=fnt,
            fill=fill,
            spacing=line_gap * self.scale,
            anchor=anchor,
            align=align,
        )

    def measure(self, value: str, size: int, bold: bool = False) -> tuple[int, int]:
        bbox = self.draw.textbbox((0, 0), value, font=font(size, bold=bold, scale=self.scale))
        return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(value: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> str:
    lines: list[str] = []
    for raw_line in value.split("\n"):
        words = raw_line.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            trial = f"{current} {word}"
            if fnt.getlength(trial) <= max_width:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return "\n".join(lines)


def section(p: Poster, x: int, y: int, w: int, h: int, title: str, accent: str = UQ_PURPLE):
    p.rect(x, y, w, h, "#FFFFFF", LINE, width=1, radius=7)
    p.rect(x, y, w, 28, accent, None, radius=7)
    p.text(title.upper(), x + 12, y + 7, 10, "#FFFFFF", bold=True)


def bullet_list(p: Poster, items: Iterable[str], x: int, y: int, w: int, size: int = 12, color: str = INK, gap: int = 31):
    yy = y
    for item in items:
        p.text("-", x, yy, size, fill=color, bold=True)
        p.text(item, x + 14, yy, size, fill=color, max_width=w - 14, line_gap=2)
        yy += gap
    return yy


def draw_metric(p: Poster, x: int, y: int, value: str, label: str, color: str, note: str = ""):
    p.rect(x, y, 101, 60, "#FFFFFF", color, width=2, radius=8)
    p.text(value, x + 50, y + 9, 20, color, bold=True, black=True, anchor="ma", align="center")
    p.text(label, x + 50, y + 34, 7, MUTED, bold=True, anchor="ma", align="center", max_width=90)
    if note:
        p.text(note, x + 50, y + 48, 6, MUTED, anchor="ma", align="center", max_width=90)


def draw_pipeline(p: Poster, x: int, y: int, w: int):
    stages = [
        ("API docs", "human intent", UQ_PURPLE),
        ("JSON specs", "LLM + validator", TEAL),
        ("Seeds", "smoke-tested", BLUE),
        ("Fuzz run", "mutations + devices", CORAL),
        ("Evidence", "pytest + reports", GREEN),
    ]
    box_w = 88
    gap = (w - box_w * len(stages)) / (len(stages) - 1)
    for idx, (title, note, color) in enumerate(stages):
        xx = x + idx * (box_w + gap)
        p.rect(xx, y, box_w, 61, "#F8FAFC", color, width=2, radius=8)
        p.text(title, xx + box_w / 2, y + 10, 10, color, bold=True, anchor="ma", align="center", max_width=80)
        p.text(note, xx + box_w / 2, y + 32, 7, MUTED, anchor="ma", align="center", max_width=76)
        if idx < len(stages) - 1:
            p.line(xx + box_w + 5, y + 30, xx + box_w + gap - 5, y + 30, "#9AA4B2", width=2)
            p.text(">", xx + box_w + gap / 2 - 3, y + 21, 15, "#9AA4B2", bold=True)


def draw_funnel(p: Poster, metrics: dict, x: int, y: int, w: int):
    totals = metrics["totals"]
    stages = [
        ("Accepted docs", totals["accepted"], UQ_PURPLE),
        ("Valid specs", totals["stage2_valid"], TEAL),
        ("Ready seeds", totals["stage3_ready"], BLUE),
        ("Evaluated", totals["stage4_eval"], CORAL),
        ("Base-valid", totals["base_valid"], GREEN),
    ]
    max_v = stages[0][1]
    yy = y
    for label, value, color in stages:
        bar_w = int((w - 150) * value / max_v)
        p.text(label, x, yy, 9, MUTED, bold=True)
        p.rect(x + 105, yy + 1, w - 152, 16, "#E8EDF3", None, radius=4)
        p.rect(x + 105, yy + 1, bar_w, 16, color, None, radius=4)
        p.text(fmt_int(value), x + w - 38, yy - 1, 9, INK, bold=True, anchor="ra")
        yy += 28


def draw_coverage(p: Poster, rows: list[dict], x: int, y: int, w: int, h: int):
    top = y + 16
    max_v = 100
    group_h = h / len(rows)
    for idx, row in enumerate(rows):
        yy = top + idx * group_h
        p.text(row["label"], x, yy, 10, INK, bold=True)
        p.text(f"v{row['version']}", x, yy + 14, 7, MUTED)
        bar_x = x + 92
        p.rect(bar_x, yy, w - 112, 10, "#E8EDF3", None, radius=4)
        p.rect(bar_x, yy, (w - 112) * row["api_exec"] / max_v, 10, TEAL, None, radius=4)
        p.text(fmt_pct(row["api_exec"], 1), x + w - 4, yy - 3, 8, TEAL, bold=True, anchor="ra")
        p.rect(bar_x, yy + 18, w - 112, 10, "#E8EDF3", None, radius=4)
        p.rect(bar_x, yy + 18, (w - 112) * row["function_cov"] / max_v, 10, CORAL, None, radius=4)
        p.text(fmt_pct(row["function_cov"], 1), x + w - 4, yy + 15, 8, CORAL, bold=True, anchor="ra")
    p.rect(x + 92, y, 10, 8, TEAL, None, radius=2)
    p.text("API execution", x + 106, y - 2, 7, MUTED)
    p.rect(x + 195, y, 10, 8, CORAL, None, radius=2)
    p.text("selected-function lines", x + 209, y - 2, 7, MUTED)


def draw_triage(p: Poster, metrics: dict, x: int, y: int, w: int):
    totals = metrics["totals"]
    items = [
        ("Candidate bugs", totals["candidate_bugs"], CORAL),
        ("Device mismatches", totals["device_mismatches"], UQ_PURPLE),
        ("Expected rejections", totals["expected_negative"], TEAL),
        ("Pipeline errors", totals["pipeline_errors"], GOLD),
        ("Platform exclusions", totals["platform_exclusions"], MUTED),
    ]
    max_log = max(math.log10(v + 1) for _, v, _ in items)
    yy = y
    for label, value, color in items:
        p.text(label, x, yy, 8, MUTED, bold=True)
        p.rect(x + 112, yy + 1, w - 160, 13, "#E8EDF3", None, radius=4)
        bw = (w - 160) * math.log10(value + 1) / max_log
        p.rect(x + 112, yy + 1, bw, 13, color, None, radius=4)
        p.text(fmt_int(value), x + w - 4, yy - 1, 8, INK, bold=True, anchor="ra")
        yy += 24
    p.text("Log-scaled bars keep the large PyTorch device-oracle count readable.", x, y + 124, 7, MUTED, max_width=w)


def draw_prior_work_table(p: Poster, x: int, y: int, w: int):
    rows = [
        ("DocTer / VISTAFUZZ", "Docs -> constraints", "DeepFuzz adds runnable artefact trail"),
        ("TitanFuzz / Fuzz4All", "LLM as input generator", "DeepFuzz validates specs before fuzzing"),
        ("XAMT", "Cross-framework oracle", "DeepFuzz focuses docs-to-test pipeline"),
        ("FlashFuzz", "Coverage-guided harnesses", "Future baseline for CGF extension"),
    ]
    col = [116, 132, w - 264]
    p.text("Related work", x, y, 9, INK, bold=True)
    p.text("Focus", x + col[0], y, 9, INK, bold=True)
    p.text("Where this work differs", x + col[0] + col[1], y, 9, INK, bold=True)
    p.line(x, y + 16, x + w, y + 16, LINE, width=1)
    yy = y + 24
    for idx, (name, focus, diff) in enumerate(rows):
        fill = "#FFFFFF" if idx % 2 == 0 else "#F8FAFC"
        p.rect(x - 3, yy - 4, w + 6, 32, fill, None, radius=3)
        p.text(name, x, yy, 7, INK, bold=True, max_width=col[0] - 6)
        p.text(focus, x + col[0], yy, 7, MUTED, max_width=col[1] - 6)
        p.text(diff, x + col[0] + col[1], yy, 7, MUTED, max_width=col[2])
        yy += 34


def create_poster(metrics: dict, background_path: Path | None) -> tuple[Path, Path, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    if background_path and background_path.exists():
        bg = Image.open(background_path).convert("RGB")
    else:
        bg = Image.new("RGB", (1587, 1123), "#FFFFFF")
    p = Poster(bg, scale=3)

    p.rect(0, 0, 1587, 1018, "#FFFFFF", None, radius=0)
    p.text("DeepFuzz: Documentation-Driven Fuzz Testing", 45, 24, 31, UQ_PURPLE, bold=True, black=True)
    p.text(
        "For deep-learning APIs: turning documentation into executable seeds, generated tests, coverage evidence, and conservative bug triage",
        47,
        63,
        13,
        INK,
        max_width=950,
    )
    p.text("Aryan Somesh Gupta | Supervisor: Dr Guowei Yang | School of Electrical Engineering and Computer Science, UQ", 47, 88, 10, MUTED)

    p.rect(1140, 25, 390, 83, "#F7F2FA", UQ_PURPLE, width=1, radius=10)
    p.text("HOOK FOR NON-SPECIALISTS", 1158, 38, 8, UQ_PURPLE, bold=True)
    p.text("$2.41T", 1158, 54, 19, CORAL, bold=True, black=True)
    p.text("cost of poor software quality in the US (CISQ 2022)", 1247, 58, 8, MUTED, max_width=250)
    p.text("46%", 1158, 82, 17, TEAL, bold=True, black=True)
    p.text("of Stack Overflow 2025 respondents distrust AI output accuracy", 1225, 84, 8, MUTED, max_width=285)
    p.line(45, 118, 1538, 118, UQ_PURPLE, width=2)

    # Left column.
    section(p, 45, 138, 450, 202, "Problem + Motivation", UQ_PURPLE)
    p.text("API docs are promises developers build on.", 60, 178, 15, INK, bold=True, max_width=415)
    bullet_list(
        p,
        [
            "Deep-learning APIs change fast and combine shapes, dtypes, devices, masks, and edge numeric values.",
            "Manual tests and examples cover the happy path, but developers lose time debugging unclear API behavior.",
            "Fuzzing means automatically trying many inputs; here the hard part is making those inputs valid and meaningful.",
        ],
        62,
        207,
        405,
        size=10,
        gap=33,
    )

    section(p, 45, 358, 450, 228, "Aim + Research Questions", TEAL)
    p.text("End goal: a working proof-of-concept plus exploratory evaluation.", 60, 398, 12, INK, bold=True, max_width=410)
    rq_items = [
        ("RQ1", "Can documentation become executable API specs and seeds?", f"Yes: {fmt_pct(metrics['totals']['stage2_rate'], 1)} valid specs."),
        ("RQ2", "Can it scale across JAX, TensorFlow, and PyTorch?", f"Yes: {fmt_int(metrics['totals']['base_valid'])} base-valid APIs."),
        ("RQ3", "Can oracles expose suspicious behavior without mixing in noise?", f"Yes, with conservative triage: {fmt_int(metrics['totals']['candidate_bugs'])} candidates."),
    ]
    yy = 432
    for tag, question, answer in rq_items:
        p.rect(61, yy - 2, 41, 28, "#E6F6F8", TEAL, width=1, radius=5)
        p.text(tag, 82, yy + 5, 9, TEAL, bold=True, anchor="ma")
        p.text(question, 113, yy - 1, 9, INK, bold=True, max_width=350)
        p.text(answer, 113, yy + 17, 8, MUTED, max_width=350)
        yy += 48

    section(p, 45, 604, 450, 286, "Background: Comparison", GOLD)
    draw_prior_work_table(p, 62, 644, 416)
    p.text(
        "Not a direct leaderboard: these systems use different API sets, oracles, environments, and bug-confirmation rules.",
        63,
        818,
        8,
        MUTED,
        max_width=410,
    )
    p.text(
        "DeepFuzz's strength is auditability: documentation -> spec -> seed -> generated pytest -> coverage/triage.",
        63,
        845,
        9,
        INK,
        bold=True,
        max_width=410,
    )

    # Middle column.
    section(p, 520, 138, 520, 282, "Method: Documentation -> Evidence", UQ_PURPLE)
    p.text("Every stage leaves an artefact that can be inspected, replayed, or challenged.", 538, 178, 12, INK, bold=True, max_width=485)
    draw_pipeline(p, 540, 220, 477)
    p.text("Stage outputs", 538, 306, 9, INK, bold=True)
    bullet_list(
        p,
        [
            "Stage 1 extracts structured JSON specs from documentation using a local LLM.",
            "Stage 2 validates and repairs specs against schemas and runtime signatures.",
            "Stage 3 materializes deterministic .init.json seeds and smoke-tests them.",
            "Stage 4 mutates inputs in isolated workers, measures coverage, and triages outcomes.",
        ],
        542,
        330,
        460,
        size=8,
        gap=20,
    )

    section(p, 520, 438, 520, 250, "Results Funnel", TEAL)
    draw_funnel(p, metrics, 542, 482, 470)
    p.text(
        "Figure 1. Documentation-derived artefact funnel across the final thesis runs. Biggest loss is seed readiness, especially for complex TensorFlow/PyTorch APIs.",
        542,
        634,
        8,
        MUTED,
        max_width=468,
    )

    section(p, 520, 706, 520, 184, "Live Demonstration", BLUE)
    bullet_list(
        p,
        [
            "Show one seed JSON generated from documentation.",
            "Open the generated pytest replay file for the same API.",
            "Run one fast replay test, then inspect coverage and triage reports.",
            "Connect the demo path back to the poster: docs -> tests -> evidence.",
        ],
        540,
        748,
        462,
        size=9,
        gap=24,
    )
    p.text("Backup: use saved artefacts if full Stage 4 is too slow for a live room.", 540, 858, 8, MUTED, max_width=470)

    # Right column.
    section(p, 1065, 138, 475, 286, "What Was Produced", CORAL)
    metric_x = [1085, 1194, 1303, 1412]
    draw_metric(p, metric_x[0], 176, fmt_int(metrics["totals"]["accepted"]), "accepted documented APIs", UQ_PURPLE)
    draw_metric(p, metric_x[1], 176, fmt_int(metrics["totals"]["generated_tests"]), "generated pytest files", TEAL)
    draw_metric(p, metric_x[2], 176, fmt_int(metrics["totals"]["valid_programs"]), "valid generated programs", BLUE)
    draw_metric(p, metric_x[3], 176, fmt_int(metrics["totals"]["candidate_bugs"]), "candidate bug records", CORAL)
    draw_coverage(p, metrics["rows"], 1086, 269, 430, 118)
    p.text(
        "Figure 2. API execution stayed near complete for evaluated APIs; selected-function line coverage exposes where deeper branch exploration remains.",
        1086,
        395,
        7,
        MUTED,
        max_width=430,
    )

    section(p, 1065, 442, 475, 214, "Triage: So What?", UQ_PURPLE)
    draw_triage(p, metrics, 1086, 482, 430)
    p.text(
        "Candidate records are not claimed as confirmed upstream bugs. They are reproducible review targets with paths to replay evidence.",
        1086,
        624,
        8,
        INK,
        bold=True,
        max_width=430,
    )

    section(p, 1065, 674, 475, 216, "Conclusion + Critical Review", GREEN)
    p.text("Main finding", 1086, 715, 9, GREEN, bold=True)
    p.text(
        "Documentation is useful test material when it is grounded by validation, seed execution, isolated fuzzing, and transparent triage.",
        1170,
        715,
        9,
        INK,
        max_width=330,
    )
    p.text("Trustworthy because", 1086, 763, 9, TEAL, bold=True)
    p.text("fixed seeds, run manifests, generated pytest files, coverage JSON, and CSV bug reports can be replayed.", 1195, 763, 8, INK, max_width=300)
    p.text("Limits", 1086, 808, 9, CORAL, bold=True)
    p.text("one-call APIs only; native/kernel coverage unavailable from binary wheels; candidates need manual confirmation.", 1135, 808, 8, INK, max_width=365)
    p.text("Next work", 1086, 850, 9, GOLD, bold=True)
    p.text("coverage-guided harnesses, cross-framework oracles, better relational constraints, and upstream confirmation workflow.", 1150, 850, 8, INK, max_width=350)

    # Reference strip.
    p.rect(45, 910, 1495, 86, "#F8FAFC", LINE, width=1, radius=7)
    p.text("References for further reading", 60, 927, 9, UQ_PURPLE, bold=True)
    refs = (
        "[1] DocTer ISSTA'22  [2] TitanFuzz ISSTA'23  [3] Fuzz4All ICSE'24  "
        "[4] VISTAFUZZ 2025  [5] XAMT 2025  [6] FlashFuzz 2025  "
        "[7] CISQ 2022 CPSQ  [8] Stack Overflow Developer Survey 2025"
    )
    p.text(refs, 60, 950, 8, MUTED, max_width=1458)
    p.text(
        "Acknowledgements: Dr Guowei Yang, Bin Duan, UQ EECS, and the maintainers of JAX, TensorFlow, PyTorch, Python, pytest, coverage.py, and local open-source LLM tooling.",
        60,
        973,
        7,
        MUTED,
        max_width=1458,
    )

    poster_png = OUT_DIR / "DeepFuzz_Demo_Day_Poster.png"
    poster_pdf = OUT_DIR / "DeepFuzz_Demo_Day_Poster.pdf"
    poster_preview = OUT_DIR / "DeepFuzz_Demo_Day_Poster_preview.png"
    p.image.save(poster_png, optimize=True)
    p.image.resize((p.logical_w, p.logical_h), Image.Resampling.LANCZOS).save(poster_preview, optimize=True)
    p.image.save(poster_pdf, "PDF", resolution=300.0)
    return poster_png, poster_pdf, poster_preview


def create_chart_assets(metrics: dict) -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    # Coverage chart.
    img = Image.new("RGB", (1400, 620), "#FFFFFF")
    d = ImageDraw.Draw(img)
    f_title = font(34, bold=True)
    f_axis = font(22, bold=True)
    f_small = font(18)
    d.text((40, 30), "DeepFuzz coverage by library", font=f_title, fill=INK)
    d.rectangle((40, 92, 62, 112), fill=TEAL)
    d.text((70, 88), "API execution coverage", font=f_small, fill=MUTED)
    d.rectangle((340, 92, 362, 112), fill=CORAL)
    d.text((370, 88), "Selected-function Python line coverage", font=f_small, fill=MUTED)
    x0, y0, w = 280, 165, 950
    for idx, row in enumerate(metrics["rows"]):
        yy = y0 + idx * 125
        d.text((40, yy + 12), f"{row['label']} {row['version']}", font=f_axis, fill=INK)
        d.rectangle((x0, yy, x0 + w, yy + 28), fill="#E8EDF3")
        d.rectangle((x0, yy, x0 + int(w * row["api_exec"] / 100), yy + 28), fill=TEAL)
        d.text((x0 + w + 15, yy - 2), fmt_pct(row["api_exec"], 1), font=f_small, fill=TEAL)
        d.rectangle((x0, yy + 45, x0 + w, yy + 73), fill="#E8EDF3")
        d.rectangle((x0, yy + 45, x0 + int(w * row["function_cov"] / 100), yy + 73), fill=CORAL)
        d.text((x0 + w + 15, yy + 43), fmt_pct(row["function_cov"], 1), font=f_small, fill=CORAL)
    d.text((40, 555), "Caption: selected-function coverage is the headline metric because whole-library native/kernel coverage was unavailable from binary wheels.", font=f_small, fill=MUTED)
    path = ASSET_DIR / "coverage_by_library.png"
    img.save(path, optimize=True)
    paths["coverage"] = path

    # Funnel chart.
    img = Image.new("RGB", (1400, 560), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((40, 30), "Final thesis artefact funnel", font=f_title, fill=INK)
    stages = [
        ("Accepted docs", metrics["totals"]["accepted"], UQ_PURPLE),
        ("Valid specs", metrics["totals"]["stage2_valid"], TEAL),
        ("Ready/evaluated", metrics["totals"]["stage4_eval"], BLUE),
        ("Base-valid", metrics["totals"]["base_valid"], GREEN),
        ("Generated tests", metrics["totals"]["generated_tests"], CORAL),
    ]
    max_v = metrics["totals"]["accepted"]
    y = 125
    for label, value, color in stages:
        d.text((60, y + 6), label, font=f_axis, fill=INK)
        d.rectangle((300, y, 1160, y + 38), fill="#E8EDF3")
        d.rectangle((300, y, 300 + int(860 * value / max_v), y + 38), fill=color)
        d.text((1185, y + 2), fmt_int(value), font=f_axis, fill=color)
        y += 75
    d.text((40, 500), f"Caption: {fmt_int(metrics['totals']['valid_programs'])} valid generated programs were produced from these evaluated APIs.", font=f_small, fill=MUTED)
    path = ASSET_DIR / "artefact_funnel.png"
    img.save(path, optimize=True)
    paths["funnel"] = path

    # Triage chart.
    img = Image.new("RGB", (1400, 560), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((40, 30), "Triage categories from final runs", font=f_title, fill=INK)
    items = [
        ("Candidate implementation-bug records", metrics["totals"]["candidate_bugs"], CORAL),
        ("Device-oracle mismatch records", metrics["totals"]["device_mismatches"], UQ_PURPLE),
        ("Expected negative rejections", metrics["totals"]["expected_negative"], TEAL),
        ("Pipeline/runtime error records", metrics["totals"]["pipeline_errors"], GOLD),
        ("Platform/runtime exclusions", metrics["totals"]["platform_exclusions"], MUTED),
    ]
    max_v = max(value for _, value, _ in items)
    y = 122
    for label, value, color in items:
        d.text((60, y + 3), label, font=f_small, fill=INK)
        d.rectangle((430, y, 1140, y + 32), fill="#E8EDF3")
        d.rectangle((430, y, 430 + int(710 * value / max_v), y + 32), fill=color)
        d.text((1170, y), fmt_int(value), font=f_axis, fill=color)
        y += 72
    d.text((40, 500), "Caption: candidate records are reproducible review targets, not automatically confirmed upstream bugs.", font=f_small, fill=MUTED)
    path = ASSET_DIR / "triage_counts.png"
    img.save(path, optimize=True)
    paths["triage"] = path
    return paths


def build_minimal_pptx(poster_png: Path, out_pptx: Path):
    width_emu = 15119350
    height_emu = 10691813
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    image_name = "poster.png"
    files: dict[str, bytes | str] = {
        "[Content_Types].xml": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Default Extension="png" ContentType="image/png"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
  <Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
  <Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>
  <Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>
  <Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>""",
        "_rels/.rels": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>""",
        "docProps/core.xml": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:title>DeepFuzz Demo Day Poster</dc:title>
  <dc:creator>Aryan Somesh Gupta</dc:creator>
  <cp:lastModifiedBy>Codex</cp:lastModifiedBy>
  <dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>
  <dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>
</cp:coreProperties>""",
        "docProps/app.xml": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
  <Application>Microsoft PowerPoint</Application>
  <PresentationFormat>A3 landscape poster</PresentationFormat>
  <Slides>1</Slides>
</Properties>""",
        "ppt/presentation.xml": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>
  <p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst>
  <p:sldSz cx="{width_emu}" cy="{height_emu}" type="custom"/>
  <p:notesSz cx="6858000" cy="9144000"/>
</p:presentation>""",
        "ppt/_rels/presentation.xml.rels": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="slideMasters/slideMaster1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
</Relationships>""",
        "ppt/slides/slide1.xml": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:bg><p:bgPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></p:bgPr></p:bg>
    <p:spTree>
      <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
      <p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>
      <p:pic>
        <p:nvPicPr><p:cNvPr id="2" name="DeepFuzz Demo Day Poster" descr="A3 poster showing DeepFuzz method, research questions, results, triage, and references."/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>
        <p:blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
        <p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{width_emu}" cy="{height_emu}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
      </p:pic>
    </p:spTree>
  </p:cSld>
  <p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>""",
        "ppt/slides/_rels/slide1.xml.rels": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/{image_name}"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>
</Relationships>""",
        "ppt/slideMasters/slideMaster1.xml": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sldMaster xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:spTree></p:cSld>
  <p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>
  <p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst>
</p:sldMaster>""",
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="../theme/theme1.xml"/>
</Relationships>""",
        "ppt/slideLayouts/slideLayout1.xml": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sldLayout xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" type="blank" preserve="1">
  <p:cSld name="Blank"><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:spTree></p:cSld>
  <p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sldLayout>""",
        "ppt/slideLayouts/_rels/slideLayout1.xml.rels": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="../slideMasters/slideMaster1.xml"/>
</Relationships>""",
        "ppt/theme/theme1.xml": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="DeepFuzz">
  <a:themeElements>
    <a:clrScheme name="DeepFuzz"><a:dk1><a:srgbClr val="14213D"/></a:dk1><a:lt1><a:srgbClr val="FFFFFF"/></a:lt1><a:dk2><a:srgbClr val="51247A"/></a:dk2><a:lt2><a:srgbClr val="F8FAFC"/></a:lt2><a:accent1><a:srgbClr val="51247A"/></a:accent1><a:accent2><a:srgbClr val="007C89"/></a:accent2><a:accent3><a:srgbClr val="D1442E"/></a:accent3><a:accent4><a:srgbClr val="2E7D32"/></a:accent4><a:accent5><a:srgbClr val="A56B00"/></a:accent5><a:accent6><a:srgbClr val="2F6FDB"/></a:accent6><a:hlink><a:srgbClr val="51247A"/></a:hlink><a:folHlink><a:srgbClr val="51247A"/></a:folHlink></a:clrScheme>
    <a:fontScheme name="Arial"><a:majorFont><a:latin typeface="Arial"/></a:majorFont><a:minorFont><a:latin typeface="Arial"/></a:minorFont></a:fontScheme>
    <a:fmtScheme name="Default"><a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst><a:lnStyleLst><a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst><a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst><a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst></a:fmtScheme>
  </a:themeElements>
</a:theme>""",
    }
    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out_pptx, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
        archive.writestr(f"ppt/media/{image_name}", poster_png.read_bytes())


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_source_list(doc: Document):
    doc.add_heading("References And Further Reading", level=1)
    refs = [
        ("DocTer", "https://arxiv.org/abs/2109.01002"),
        ("TitanFuzz", "https://arxiv.org/abs/2212.14834"),
        ("Fuzz4All", "https://arxiv.org/abs/2308.04748"),
        ("VISTAFUZZ", "https://arxiv.org/abs/2507.14558"),
        ("XAMT", "https://arxiv.org/abs/2508.12546"),
        ("FlashFuzz", "https://arxiv.org/abs/2509.14626"),
        ("CISQ Cost of Poor Software Quality 2022", "https://www.it-cisq.org/press-releases/12-06-22/"),
        ("Stack Overflow Developer Survey 2025 AI section", "https://survey.stackoverflow.co/2025/ai"),
    ]
    for label, url in refs:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{label}: ").bold = True
        p.add_run(url)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, UQ_PURPLE)
        set_cell_text(cell, header, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            if idx == 0:
                cells[idx].paragraphs[0].runs[0].bold = True
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def create_docx(metrics: dict, chart_paths: dict[str, Path], poster_preview: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for name, size, color in [
        ("Title", 24, UQ_PURPLE),
        ("Heading 1", 16, UQ_PURPLE),
        ("Heading 2", 13, TEAL),
        ("Heading 3", 11, INK),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
        style.font.bold = name != "Title"

    title = doc.add_paragraph()
    title.style = styles["Title"]
    title.add_run("DeepFuzz Demo Day Helper Pack")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Poster script, Q&A preparation, demo flow, supplementary figures, and references").italic = True
    subtitle.add_run("\nAryan Somesh Gupta | Supervisor: Dr Guowei Yang | UQ EECS")

    doc.add_heading("One-Page Brief", level=1)
    p = doc.add_paragraph()
    p.add_run("Project type: ").bold = True
    p.add_run("Completed proof-of-concept and exploratory evaluation. DeepFuzz is working research software with reproducible artefacts, not a claim of production readiness.")
    p = doc.add_paragraph()
    p.add_run("Main message: ").bold = True
    p.add_run("API documentation can become executable fuzzing evidence when it is validated, smoke-tested, mutated, measured, and triaged conservatively.")

    add_table(
        doc,
        ["Metric", "Final result", "Meaning"],
        [
            ["Accepted documented APIs", fmt_int(metrics["totals"]["accepted"]), "Input scope across JAX, TensorFlow, and PyTorch."],
            ["Valid structured specs", fmt_int(metrics["totals"]["stage2_valid"]), "Documentation-to-JSON conversion mostly survived validation."],
            ["Base-valid APIs", fmt_int(metrics["totals"]["base_valid"]), "Seeds executed successfully in Stage 4."],
            ["Generated pytest files", fmt_int(metrics["totals"]["generated_tests"]), "Replayable test artefacts for demo and regression use."],
            ["Valid generated programs", fmt_int(metrics["totals"]["valid_programs"]), "Executed base/mutated API calls."],
            ["Candidate bug records", fmt_int(metrics["totals"]["candidate_bugs"]), "Review targets, not automatically confirmed upstream bugs."],
        ],
        [1.7, 1.2, 3.5],
    )

    doc.add_heading("30-Second Elevator Pitch", level=1)
    pitch = (
        "Every developer has lost time asking: is my code wrong, or did the library/API behave differently than promised? "
        "DeepFuzz tackles that pain for deep-learning libraries. It turns API documentation into executable fuzz tests, "
        "runs them across JAX, TensorFlow and PyTorch, and produces replayable evidence: seeds, pytest files, coverage reports and candidate bug triage."
    )
    doc.add_paragraph(pitch)

    doc.add_heading("Five-Minute Poster Walk-Through", level=1)
    script_rows = [
        ["0:00-0:30", "Hook", "Start with developer toil: debugging library/API behavior wastes time because documentation is human-readable but not executable."],
        ["0:30-1:05", "Aim and RQs", "State the end goal: a reproducible proof-of-concept pipeline. Point to RQ1-RQ3 and say what evidence answers each one."],
        ["1:05-1:55", "Background", "Explain fuzzing as automated input exploration. Compare DeepFuzz to DocTer, TitanFuzz, Fuzz4All, XAMT, VISTAFUZZ, and FlashFuzz without claiming a direct leaderboard."],
        ["1:55-2:55", "Method", "Walk left to right through docs -> JSON specs -> seeds -> isolated fuzzing -> pytest/coverage/triage. Emphasise that every step leaves files."],
        ["2:55-3:55", "Results", "Use the funnel, coverage chart, and metrics. Explain TensorFlow's lower selected-function coverage as complex wrappers/eager behavior and harder seed constraints."],
        ["3:55-4:35", "Triage", "Say candidate bugs are conservative review targets. Device mismatches show where CPU/accelerator behavior deserves investigation."],
        ["4:35-5:00", "Close", "DeepFuzz turns documentation into evidence. It is trustworthy because results are replayable; future work is deeper constraints, coverage-guided harnesses, and upstream confirmation."],
    ]
    add_table(doc, ["Time", "Beat", "What to say"], script_rows, [0.9, 1.25, 4.25])

    doc.add_heading("Interaction Prompts", level=2)
    for item in [
        "Ask: Have you ever had code fail and wondered whether the library, the docs, or your code was wrong?",
        "Ask: If I can show the exact seed and generated test that produced a result, would you trust it more?",
        "Ask technical listeners whether they care more about coverage, bug triage, or reproducibility, then route the demo accordingly.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Live Demo Plan", level=1)
    doc.add_paragraph(
        "Run commands from /Users/aryansg/Desktop/DeepFuzz. Activate the thesis Python environment first; replay tests need the target library "
        "installed (JAX, TensorFlow, or PyTorch). Use the fast path first, then fall back to saved final artefacts if time is tight or the room machine "
        "does not have the deep-learning stack installed."
    )
    demo_commands = [
        ("Show the latest mini demo summary", "python3 -m json.tool demo-run/results/latest/demo_report.json | sed -n '1,120p'"),
        ("Inspect one documentation-derived seed", "sed -n '1,120p' json2init/results/jax/jax.numpy.sin.init.json"),
        ("Inspect the replayable pytest", "sed -n '1,120p' pipeline_runs/jax/jax-thesis-all/generated_tests/test_jax_numpy_sin.py"),
        ("Run one quick replay test, after env activation", "python3 -m pytest pipeline_runs/jax/jax-thesis-all/generated_tests/test_jax_numpy_sin.py -q"),
        ("Show final PyTorch evidence", "python3 scripts/read_coverage.py --stage4-results-dir stage4/results/torch-thesis-all-coverage"),
        ("Open candidate triage evidence", "sed -n '1,12p' stage4/results/torch-thesis-all-coverage/bug_report.csv"),
    ]
    add_table(doc, ["Demo step", "Command"], [[a, b] for a, b in demo_commands], [2.2, 4.4])

    doc.add_heading("Supplementary Figures", level=1)
    for key, caption in [
        ("funnel", "Figure A. Final thesis artefact funnel. The largest loss is seed readiness, showing where better relational constraints matter."),
        ("coverage", "Figure B. Coverage by library. API execution is high for evaluated APIs, while selected-function coverage shows room for deeper branch exploration."),
        ("triage", "Figure C. Triage categories. Candidate bug records are separated from expected negative tests, pipeline errors, and platform exclusions."),
    ]:
        doc.add_picture(str(chart_paths[key]), width=Inches(6.35))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Practice Questions And Strong Answers", level=1)
    qa = [
        (
            "What is the end goal of the thesis?",
            "A completed proof-of-concept and exploratory evaluation: DeepFuzz demonstrates an end-to-end, reproducible path from documentation to executable fuzzing artefacts and evidence. It is not claiming production deployment or confirmed upstream bug totals.",
        ),
        (
            "Why choose this topic?",
            "Developer toil. When APIs behave unexpectedly, people waste time deciding whether the bug is in their code, the documentation, or the library. Deep-learning APIs make this worse because valid inputs depend on shapes, dtypes, devices, masks, and numerical edge cases.",
        ),
        (
            "Why fuzzing over normal unit tests?",
            "Unit tests are great for known examples. Fuzzing explores more of the input space and can expose edge cases a developer did not manually imagine. DeepFuzz makes fuzzing practical by generating valid seeds from documentation first.",
        ),
        (
            "Why use documentation if documentation can be wrong?",
            "That is exactly why the pipeline validates. Documentation is treated as an intent source, then checked through schemas, runtime signatures, smoke tests, isolated execution, and triage. Mismatch becomes evidence, not blind trust.",
        ),
        (
            "What did you solve from the RQs?",
            "RQ1: documentation-to-spec validation mostly worked. RQ2: the pipeline generated replayable tests and valid programs across three libraries. RQ3: triage produced candidate implementation records while separating expected negative inputs and pipeline failures.",
        ),
        (
            "How does this compare with TitanFuzz or Fuzz4All?",
            "TitanFuzz and Fuzz4All show LLMs can generate fuzz inputs broadly. DeepFuzz is more documentation-to-evidence oriented: it preserves structured specs, deterministic seeds, generated pytest files, coverage reports, and conservative CSV/JSON triage.",
        ),
        (
            "How does this compare with DocTer?",
            "DocTer extracts documentation constraints with rule-based NLP and reports confirmed bugs/document inconsistencies. DeepFuzz keeps the documentation-first idea but adds LLM-assisted structured specs, deterministic seed files, generated pytest replays, and selected-function coverage.",
        ),
        (
            "Are your 190 candidates confirmed bugs?",
            "No. They are candidate implementation-bug records. The trustworthy claim is that they are reproducible review targets, not that all have been accepted upstream.",
        ),
        (
            "Why is TensorFlow selected-function coverage lower?",
            "TensorFlow has many wrappers, eager/graph boundaries, complex argument constraints, and APIs where a one-call seed reaches the public wrapper but not many internal branches. That is a limitation and a future-work target.",
        ),
        (
            "Why is native/kernel coverage unavailable?",
            "The final environment used installed binary wheels. Native coverage requires instrumented source builds and backend/kernel coverage tooling. Reporting it as unavailable is more honest than inventing a whole-library number.",
        ),
        (
            "Why should findings still be trusted?",
            "The pipeline uses fixed seeds, run manifests, generated pytest files, coverage JSON, and triage CSV/JSON. Even if a candidate is later classified as not-a-bug, the evidence path is reproducible and auditable.",
        ),
        (
            "What should come next?",
            "Add coverage-guided harness generation, relational constraint learning, cross-framework oracles, native/kernel coverage with source builds, and an upstream reporting workflow that confirms or rejects candidate records.",
        ),
    ]
    for question, answer in qa:
        p = doc.add_paragraph()
        p.add_run(question).bold = True
        doc.add_paragraph(answer)

    doc.add_heading("Limitations And How To Defend Them", level=1)
    add_table(
        doc,
        ["Limitation", "Why it happened", "How to defend it"],
        [
            ["One-call APIs only", "Safe thesis scope and reproducibility.", "Clear denominator; avoids noisy multi-step object setup."],
            ["Candidate bugs unconfirmed", "Manual upstream confirmation takes time.", "Use conservative language and show replay artefacts."],
            ["No native/kernel coverage", "Binary wheels are not instrumented.", "Report selected-function Python coverage honestly."],
            ["Seed failures remain", "Docs often omit relational constraints.", "These failures are useful evidence for future constraint learning."],
            ["PyTorch mismatch volume", "Device oracle plus edge cases can create many related candidates.", "Cluster, deduplicate, and manually review before upstream claims."],
        ],
        [1.55, 2.25, 2.6],
    )

    add_source_list(doc)
    doc.add_heading("Poster Preview", level=1)
    doc.add_picture(str(poster_preview), width=Inches(6.4))
    doc.add_paragraph("Use the PDF for printing; use the PPTX if the assessment workflow requests a PowerPoint file.")

    out = OUT_DIR / "DeepFuzz_Demo_Day_Helper.docx"
    doc.save(out)
    return out


def write_summary(metrics: dict, files: dict[str, Path]):
    summary = {
        "created_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "metrics": metrics,
        "files": {key: str(value) for key, value in files.items()},
    }
    (OUT_DIR / "deliverable_manifest.json").write_text(json.dumps(summary, indent=2) + "\n")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--background", type=Path, default=None)
    args = parser.parse_args()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    metrics = collect_metrics()
    background = args.background or latest_template_background()
    poster_png, poster_pdf, poster_preview = create_poster(metrics, background)
    chart_paths = create_chart_assets(metrics)
    pptx = OUT_DIR / "DeepFuzz_Demo_Day_Poster.pptx"
    build_minimal_pptx(poster_png, pptx)
    helper_docx = create_docx(metrics, chart_paths, poster_preview)
    write_summary(
        metrics,
        {
            "poster_png": poster_png,
            "poster_pdf": poster_pdf,
            "poster_preview": poster_preview,
            "poster_pptx": pptx,
            "helper_docx": helper_docx,
            **{f"chart_{key}": value for key, value in chart_paths.items()},
        },
    )
    print(pptx)
    print(poster_pdf)
    print(helper_docx)


if __name__ == "__main__":
    main()
